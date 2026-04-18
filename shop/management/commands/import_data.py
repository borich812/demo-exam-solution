from pathlib import Path
from django.contrib.auth.models import User
from django.core.management.base import BaseCommand
from openpyxl import load_workbook
from shop.models import Category, CakeType, Manufacturer, Order, OrderItem, OrderStatus, PickupPoint, Product, Profile, Supplier


ROLE_MAP = {
    'Администратор': Profile.ROLE_ADMIN,
    'Менеджер': Profile.ROLE_MANAGER,
    'Авторизированный клиент': Profile.ROLE_CLIENT,
}


class Command(BaseCommand):
    help = 'Импорт данных из Excel-файлов задания'

    def handle(self, *args, **options):
        base_dir = Path(__file__).resolve().parents[4] / 'data' / 'import'
        self.import_users(base_dir / 'user_import.xlsx')
        self.import_products(base_dir / 'tovar.xlsx')
        self.import_pickup_points(base_dir)
        self.import_orders(base_dir)
        self.stdout.write(self.style.SUCCESS('Импорт завершен.'))

    def import_users(self, path: Path):
        ws = load_workbook(path, data_only=True).active
        for role_name, full_name, login, password in ws.iter_rows(min_row=2, values_only=True):
            if not login:
                continue
            user, _ = User.objects.get_or_create(username=login, defaults={'email': login})
            user.email = login
            user.first_name = full_name.split()[1] if len(full_name.split()) > 1 else full_name
            user.last_name = full_name.split()[0]
            user.set_password(password)
            user.save()
            profile = user.profile
            profile.full_name = full_name
            profile.role = ROLE_MAP[role_name]
            profile.save()

    def import_products(self, path: Path):
        ws = load_workbook(path, data_only=True).active
        for row in ws.iter_rows(min_row=2, values_only=True):
            article, name, unit, price, brand, cake_type, category, discount, stock, description, photo = row
            supplier, _ = Supplier.objects.get_or_create(name=brand)
            manufacturer, _ = Manufacturer.objects.get_or_create(name=brand)
            cake_type_obj, _ = CakeType.objects.get_or_create(name=cake_type)
            category_obj, _ = Category.objects.get_or_create(name=category)
            Product.objects.update_or_create(
                article=article,
                defaults={
                    'name': name,
                    'unit': unit,
                    'price': price,
                    'manufacturer': manufacturer,
                    'supplier': supplier,
                    'cake_type': cake_type_obj,
                    'category': category_obj,
                    'discount': int(discount or 0),
                    'stock': int(stock or 0),
                    'description': description or '',
                    'image': f'images/{photo}' if photo else 'images/picture.png',
                },
            )

    def import_pickup_points(self, base_dir: Path):
        pickup_path = next(base_dir.glob('*выдачи*.xlsx'))
        ws = load_workbook(pickup_path, data_only=True).active
        for (address,) in ws.iter_rows(min_row=1, values_only=True):
            PickupPoint.objects.get_or_create(address=address)

    def import_orders(self, base_dir: Path):
        order_path = next(base_dir.glob('*Заказ*.xlsx'))
        ws = load_workbook(order_path, data_only=True).active
        for row in ws.iter_rows(min_row=2, values_only=True):
            number, items_raw, order_date, delivery_date, pickup_index, customer_name, pickup_code, status_name = row
            customer = Profile.objects.get(full_name=customer_name)
            pickup_point = PickupPoint.objects.all()[int(pickup_index) - 1]
            status, _ = OrderStatus.objects.get_or_create(name=status_name)
            order, _ = Order.objects.update_or_create(
                id=int(number),
                defaults={
                    'customer': customer,
                    'pickup_point': pickup_point,
                    'status': status,
                    'order_date': order_date.date(),
                    'delivery_date': delivery_date.date(),
                    'pickup_code': str(pickup_code),
                },
            )
            order.items.all().delete()
            parts = [part.strip() for part in str(items_raw).split(',')]
            for i in range(0, len(parts), 2):
                product = Product.objects.get(article=parts[i])
                quantity = int(parts[i + 1])
                OrderItem.objects.create(order=order, product=product, quantity=quantity)
